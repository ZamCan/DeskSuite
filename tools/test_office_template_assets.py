#!/usr/bin/env python3
"""Host-side integrity/content checks for the bundled fictional templates.

Requires Python, python-docx, openpyxl and Pillow. This does NOT run Flutter,
native pickers, Office applications, layout rendering, printing or Android.
"""
from pathlib import Path
import posixpath
import unittest
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / 'assets/samples/office_templates'
EXPECTED = {
    'professional-letter.docx', 'project-brief.docx',
    'monthly-budget.xlsx', 'stock-tracker.xlsx',
    'project-pitch.pptx', 'workshop-lesson.pptx',
}


class TemplateAssetTests(unittest.TestCase):
    def test_all_archives_are_bounded_valid_xml_with_resolved_internal_links(self):
        self.assertEqual({path.name for path in ASSETS.iterdir() if path.is_file()}, EXPECTED)
        for path in sorted(ASSETS.iterdir()):
            with self.subTest(path=path.name), ZipFile(path) as package:
                self.assertIsNone(package.testzip())
                names = set(package.namelist())
                self.assertEqual(len(names), len(package.namelist()), 'duplicate ZIP names')
                self.assertLessEqual(len(names), 4000)
                self.assertLessEqual(sum(i.file_size for i in package.infolist()), 64 * 1024 * 1024)
                for info in package.infolist():
                    self.assertLessEqual(info.file_size, 24 * 1024 * 1024)
                    self.assertIn(info.compress_type, (0, 8))
                    self.assertFalse(info.flag_bits & 0x41)
                    self.assertNotEqual((info.external_attr >> 16) & 0xF000, 0xA000)
                    if info.filename.endswith(('.xml', '.rels')):
                        ET.fromstring(package.read(info.filename))
                    if info.filename.endswith('.rels'):
                        base = '' if info.filename == '_rels/.rels' else info.filename.rsplit('/_rels/', 1)[0] + '/'
                        for rel in ET.fromstring(package.read(info.filename)):
                            self.assertNotEqual(rel.get('TargetMode'), 'External', info.filename)
                            target = rel.get('Target', '')
                            resolved = posixpath.normpath(target.lstrip('/') if target.startswith('/') else base + target)
                            self.assertIn(resolved, names, f'{info.filename} → {target}')

    def test_word_templates_have_fictional_editable_text_and_metadata(self):
        for file_name, heading in [('professional-letter.docx', 'Professional letter'),
                                   ('project-brief.docx', 'Project brief')]:
            doc = Document(ASSETS / file_name)
            text = '\n'.join(p.text for p in doc.paragraphs)
            self.assertIn(heading, text)
            self.assertIn('[', text)
            self.assertEqual(doc.core_properties.author, 'DeskSuite')
            self.assertTrue(doc.sections[0].footer.paragraphs)

    def test_spreadsheets_include_real_formulas_styles_and_multiple_sheets(self):
        for file_name, expected_formula in [('monthly-budget.xlsx', '=C5-B5'),
                                            ('stock-tracker.xlsx', '=B5*C5')]:
            book = load_workbook(ASSETS / file_name, data_only=False)
            self.assertEqual(book.sheetnames, ['Template', 'Read me'])
            sheet = book['Template']
            self.assertEqual(sheet['D5'].value, expected_formula)
            self.assertEqual(sheet['D5'].data_type, 'f')
            self.assertEqual(sheet.freeze_panes, 'A5')
            self.assertTrue(sheet['A4'].font.bold)
            self.assertIn('formulas', book['Read me']['A3'].value)

    def test_presentations_have_wide_editable_shape_and_text_slides(self):
        ns = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
              'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
              'dc': 'http://purl.org/dc/elements/1.1/'}
        for file_name, count, marker in [('project-pitch.pptx', 4, 'Your project'),
                                         ('workshop-lesson.pptx', 3, 'Your workshop')]:
            with ZipFile(ASSETS / file_name) as package:
                presentation = ET.fromstring(package.read('ppt/presentation.xml'))
                self.assertEqual(len(presentation.findall('.//p:sldId', ns)), count)
                size = presentation.find('p:sldSz', ns)
                self.assertEqual((size.get('cx'), size.get('cy')), ('12192000', '6858000'))
                core = ET.fromstring(package.read('docProps/core.xml'))
                self.assertEqual(core.find('dc:creator', ns).text, 'DeskSuite')
                for index in range(1, count + 1):
                    slide = ET.fromstring(package.read(f'ppt/slides/slide{index}.xml'))
                    shapes = slide.findall('.//p:sp', ns)
                    self.assertGreaterEqual(len(shapes), 5)
                    words = ' '.join(node.text or '' for node in slide.findall('.//a:t', ns))
                    self.assertTrue(words.strip())
                    if index == 1:
                        self.assertIn(marker, words)

    def test_still_life_and_practice_pack_assets_are_present(self):
        samples = ROOT / 'assets/samples'
        with Image.open(samples / 'studio-still-life.jpg') as image:
            self.assertEqual(image.size, (1200, 896))
            self.assertEqual(image.format, 'JPEG')
            image.verify()
        for name in ['letter.png', 'invoice.png', 'application-form.png',
                     'clean-practice-page.png', 'shadowed-practice-page.png',
                     'angled-practice-page.jpg']:
            with Image.open(samples / name) as image:
                self.assertGreater(image.width, 0)
                image.verify()


if __name__ == '__main__':
    unittest.main(verbosity=2)
