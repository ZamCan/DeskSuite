#!/usr/bin/env python3
"""Generate DeskSuite's fictional, editable OOXML starter files.

Development-only requirements: python-docx and openpyxl. PPTX files are made
from the project-generated, fictional test fixture as an OOXML scaffold; slides
are replaced with original editable shapes/text and original metadata. No real
names, company documents, third-party photos or legacy-format fakes are used.
These are sample files, not proof of Office or Android rendering fidelity.
"""
from datetime import datetime
from io import BytesIO
from pathlib import Path
import re
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'assets/samples/office_templates'
SCAFFOLD = ROOT / 'test/fixtures/office_preview/preview.pptx'
STAMP = datetime(2026, 1, 1, 0, 0, 0)
NAVY = '122745'
BLUE = '245AAF'
TEAL = '3F9A91'
LIGHT = 'EAF2FC'
GRAY = '52647B'


def normalized_zip(data):
    """Repack with stable ordering/timestamps, ordinary files, CRC and deflate."""
    result = BytesIO()
    with ZipFile(BytesIO(data)) as src, ZipFile(result, 'w', ZIP_DEFLATED, compresslevel=9) as dst:
        for name in sorted(src.namelist()):
            info = ZipInfo(name, (2026, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o100644 << 16
            content = src.read(name)
            if name == 'docProps/core.xml':
                # openpyxl resets modified time at save even when properties
                # were set explicitly. Keep generated assets reproducible.
                content = re.sub(
                    rb'(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)',
                    rb'\g<1>2026-01-01T00:00:00Z\g<2>', content,
                )
            dst.writestr(info, content)
    return result.getvalue()


def add_bottom_border(paragraph, color=BLUE):
    properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    for key, value in [('val', 'single'), ('sz', '16'), ('space', '8'), ('color', color)]:
        bottom.set(qn('w:' + key), value)
    borders.append(bottom)
    properties.append(borders)


def word_template(path, title, subtitle, sections):
    doc = Document()
    doc.core_properties.title = title
    doc.core_properties.subject = 'Fictional, editable starter template'
    doc.core_properties.author = 'DeskSuite'
    doc.core_properties.last_modified_by = 'DeskSuite'
    doc.core_properties.created = STAMP
    doc.core_properties.modified = STAMP
    doc.core_properties.comments = 'Replace all bracketed placeholders before use.'
    section = doc.sections[0]
    section.top_margin = Inches(.75)
    section.bottom_margin = Inches(.7)
    section.left_margin = Inches(.85)
    section.right_margin = Inches(.85)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(8)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run('YOUR ORGANISATION  /  DESKSUITE TEMPLATE')
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(BLUE)
    add_bottom_border(p)
    p = doc.add_paragraph()
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(26)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph(subtitle)
    p.runs[0].italic = True
    p.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    for heading, body in sections:
        p = doc.add_paragraph()
        r = p.add_run(heading.upper())
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor.from_string(BLUE)
        for line in body:
            doc.add_paragraph(line)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run('FICTIONAL TEMPLATE  •  REPLACE PLACEHOLDERS')
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(GRAY)
    output = BytesIO()
    doc.save(output)
    path.write_bytes(normalized_zip(output.getvalue()))


def spreadsheet(path, *, title, subtitle, columns, rows, formula_column, notes):
    book = Workbook()
    book.properties.creator = 'DeskSuite'
    book.properties.lastModifiedBy = 'DeskSuite'
    book.properties.title = title
    book.properties.subject = 'Fictional editable starter template'
    book.properties.created = STAMP
    book.properties.modified = STAMP
    sheet = book.active
    sheet.title = 'Template'
    sheet.sheet_view.showGridLines = False
    last = get_column_letter(len(columns))
    sheet.merge_cells(f'A1:{last}1')
    sheet['A1'] = title.upper()
    sheet['A1'].font = Font(name='Calibri', size=19, bold=True, color='FFFFFF')
    sheet['A1'].fill = PatternFill('solid', fgColor=NAVY)
    sheet['A1'].alignment = Alignment(vertical='center')
    sheet.row_dimensions[1].height = 36
    sheet.merge_cells(f'A2:{last}2')
    sheet['A2'] = subtitle
    sheet['A2'].font = Font(name='Calibri', size=10, italic=True, color=GRAY)
    sheet.row_dimensions[2].height = 26
    for i, (name, width) in enumerate(columns, 1):
        col = get_column_letter(i)
        sheet.column_dimensions[col].width = width
        cell = sheet.cell(4, i, name)
        cell.font = Font(name='Calibri', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=BLUE)
        cell.alignment = Alignment(vertical='center')
    sheet.row_dimensions[4].height = 25
    thin = Side(style='thin', color='DDE6F0')
    for row_index, row in enumerate(rows, 5):
        sheet.row_dimensions[row_index].height = 24
        for col_index, value in enumerate(row, 1):
            cell = sheet.cell(row_index, col_index, value)
            cell.font = Font(name='Calibri', size=11, color=NAVY)
            cell.fill = PatternFill('solid', fgColor=('F2F7FC' if row_index % 2 == 0 else 'FFFFFF'))
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical='center')
            if col_index in formula_column:
                cell.number_format = '#,##0.00;[Red]-#,##0.00'
    sheet.freeze_panes = 'A5'
    sheet.auto_filter.ref = f'A4:{last}{4 + len(rows)}'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.print_options.horizontalCentered = True
    help_sheet = book.create_sheet('Read me')
    help_sheet.column_dimensions['A'].width = 85
    for index, text in enumerate(notes, 1):
        cell = help_sheet.cell(index, 1, text)
        cell.font = Font(name='Calibri', size=11, bold=index == 1, color=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        help_sheet.row_dimensions[index].height = 30 if index > 1 else 25
    output = BytesIO()
    book.save(output)
    path.write_bytes(normalized_zip(output.getvalue()))


# OOXML editable vector/text shapes. Coordinates are in EMUs on a wide slide.
EMU = 914400
SLIDE_W = 12192000
SLIDE_H = 6858000

def emu(inches):
    return round(inches * EMU)


def shape(shape_id, name, x, y, w, h, *, fill=None, text=None, size=18,
          color=NAVY, bold=False, rounded=False):
    fill_xml = (f'<a:solidFill><a:srgbClr val="{fill}"/></a:solidFill>'
                if fill else '<a:noFill/>')
    geometry = 'roundRect' if rounded else 'rect'
    tx = ''
    if text is not None:
        paragraphs = []
        for line in text.split('\n'):
            paragraphs.append(
                f'<a:p><a:pPr algn="l"/><a:r><a:rPr lang="en-US" sz="{size * 100}" '
                f'b="{int(bold)}"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
                f'<a:latin typeface="Aptos"/></a:rPr><a:t>{escape(line)}</a:t></a:r>'
                f'<a:endParaRPr lang="en-US"/></a:p>')
        tx = ('<p:txBody><a:bodyPr wrap="square" lIns="0" rIns="0" tIns="0" bIns="0"/>'
              '<a:lstStyle/>' + ''.join(paragraphs) + '</p:txBody>')
    return (f'<p:sp><p:nvSpPr><p:cNvPr id="{shape_id}" name="{escape(name)}"/>'
            f'<p:cNvSpPr{" txBox=" + chr(34) + "1" + chr(34) if text is not None else ""}/>'
            f'<p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="{emu(x)}" y="{emu(y)}"/>'
            f'<a:ext cx="{emu(w)}" cy="{emu(h)}"/></a:xfrm>'
            f'<a:prstGeom prst="{geometry}"><a:avLst/></a:prstGeom>{fill_xml}'
            f'<a:ln><a:noFill/></a:ln></p:spPr>{tx}</p:sp>')


def slide(background, shapes):
    head = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
            'showMasterSp="0"><p:cSld><p:bg><p:bgPr><a:solidFill>'
            f'<a:srgbClr val="{background}"/></a:solidFill><a:effectLst/></p:bgPr></p:bg>'
            '<p:spTree><p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/>'
            '<p:nvPr/></p:nvGrpSpPr><p:grpSpPr/>')
    return (head + ''.join(shapes) + '</p:spTree></p:cSld><p:clrMapOvr>'
            '<a:masterClrMapping/></p:clrMapOvr></p:sld>').encode()


def make_presentation(path, title, slides):
    with ZipFile(SCAFFOLD) as base:
        parts = {name: base.read(name) for name in base.namelist()}
    ns_p = 'http://schemas.openxmlformats.org/presentationml/2006/main'
    ns_r = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    ns_rel = 'http://schemas.openxmlformats.org/package/2006/relationships'
    ns_ct = 'http://schemas.openxmlformats.org/package/2006/content-types'
    presentation = ET.fromstring(parts['ppt/presentation.xml'])
    slist = presentation.find(f'{{{ns_p}}}sldIdLst')
    assert slist is not None
    for item in list(slist):
        slist.remove(item)
    size = presentation.find(f'{{{ns_p}}}sldSz')
    assert size is not None
    size.set('cx', str(SLIDE_W)); size.set('cy', str(SLIDE_H)); size.set('type', 'screen16x9')
    rels = ET.fromstring(parts['ppt/_rels/presentation.xml.rels'])
    for rel in list(rels):
        if rel.get('Type', '').endswith('/slide'):
            rels.remove(rel)
    types = ET.fromstring(parts['[Content_Types].xml'])
    for entry in list(types):
        if entry.get('PartName', '').startswith('/ppt/slides/slide'):
            types.remove(entry)
    for name in list(parts):
        if name.startswith('ppt/slides/slide') or name.startswith('ppt/slides/_rels/slide'):
            del parts[name]
    relation = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                f'<Relationships xmlns="{ns_rel}"><Relationship Id="rId1" '
                f'Type="{ns_r}/slideLayout" Target="../slideLayouts/slideLayout2.xml"/>'
                '</Relationships>').encode()
    for index, content in enumerate(slides, 1):
        rid = f'rId{10 + index}'
        ET.SubElement(slist, f'{{{ns_p}}}sldId', {'id': str(255 + index), f'{{{ns_r}}}id': rid})
        ET.SubElement(rels, f'{{{ns_rel}}}Relationship',
                      {'Id': rid, 'Type': ns_r + '/slide', 'Target': f'slides/slide{index}.xml'})
        ET.SubElement(types, f'{{{ns_ct}}}Override',
                      {'PartName': f'/ppt/slides/slide{index}.xml',
                       'ContentType': 'application/vnd.openxmlformats-officedocument.presentationml.slide+xml'})
        parts[f'ppt/slides/slide{index}.xml'] = content
        parts[f'ppt/slides/_rels/slide{index}.xml.rels'] = relation
    # Preserve the conventional OOXML prefixes as well as namespace URIs;
    # some real-world renderers rely on those prefixes when resolving slides.
    ET.register_namespace('p', ns_p)
    ET.register_namespace('r', ns_r)
    ET.register_namespace('a', 'http://schemas.openxmlformats.org/drawingml/2006/main')
    parts['ppt/presentation.xml'] = ET.tostring(presentation, encoding='utf-8', xml_declaration=True)
    ET.register_namespace('', ns_rel)
    parts['ppt/_rels/presentation.xml.rels'] = ET.tostring(rels, encoding='utf-8', xml_declaration=True)
    ET.register_namespace('', ns_ct)
    parts['[Content_Types].xml'] = ET.tostring(types, encoding='utf-8', xml_declaration=True)
    core = ET.fromstring(parts['docProps/core.xml'])
    for namespace, tag, value in [
        ('http://purl.org/dc/elements/1.1/', 'title', title),
        ('http://purl.org/dc/elements/1.1/', 'creator', 'DeskSuite'),
        ('http://purl.org/dc/elements/1.1/', 'description', 'Fictional editable starter template'),
        ('http://schemas.openxmlformats.org/package/2006/metadata/core-properties', 'lastModifiedBy', 'DeskSuite'),
    ]:
        node = core.find(f'{{{namespace}}}{tag}')
        if node is not None:
            node.text = value
    parts['docProps/core.xml'] = ET.tostring(core, encoding='utf-8', xml_declaration=True)
    app = ET.fromstring(parts['docProps/app.xml'])
    for node in app:
        if node.tag.endswith('}Slides'):
            node.text = str(len(slides))
    parts['docProps/app.xml'] = ET.tostring(app, encoding='utf-8', xml_declaration=True)
    output = BytesIO()
    with ZipFile(output, 'w', ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    path.write_bytes(normalized_zip(output.getvalue()))


def cover(title, subtitle, tag):
    return slide(NAVY, [
        shape(2, 'Accent panel', 0, 0, .16, 7.5, fill=BLUE),
        shape(3, 'Eyebrow', .75, .62, 10, .35, text=tag.upper(), size=12, color='85B9FF', bold=True),
        shape(4, 'Title', .75, 1.55, 10.8, 1.45, text=title, size=42, color='FFFFFF', bold=True),
        shape(5, 'Subtitle', .78, 3.35, 9.8, .8, text=subtitle, size=20, color='D3E2F7'),
        shape(6, 'Rule', .78, 5.7, 2.2, .05, fill=TEAL),
        shape(7, 'Footer', .78, 6.35, 10.5, .4, text='FICTIONAL TEMPLATE  •  REPLACE ALL PLACEHOLDERS', size=10, color='B3C6DD'),
    ])


def section_slide(title, eyebrow, cards):
    content = [
        shape(2, 'Top accent', 0, 0, 13.333, .12, fill=BLUE),
        shape(3, 'Eyebrow', .65, .52, 11, .35, text=eyebrow.upper(), size=11, color=BLUE, bold=True),
        shape(4, 'Heading', .65, 1.00, 11.8, .8, text=title, size=32, color=NAVY, bold=True),
    ]
    x = .65
    for i, (label, body) in enumerate(cards):
        content.extend([
            shape(10+i*3, 'Card background', x, 2.15, 3.83, 3.50, fill='EAF2FC', rounded=True),
            shape(11+i*3, 'Card title', x+.25, 2.50, 3.3, .55, text=label, size=18, color=BLUE, bold=True),
            shape(12+i*3, 'Card text', x+.25, 3.25, 3.25, 1.8, text=body, size=15, color=NAVY),
        ])
        x += 4.12
    content.append(shape(30, 'Footer', .65, 6.75, 10, .25,
                         text='DESKSUITE  /  FICTIONAL STARTER', size=9, color=GRAY))
    return slide('FFFFFF', content)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    word_template(OUT / 'professional-letter.docx', 'Professional letter',
                  'An editable starting point · replace all bracketed fields', [
        ('Details', ['[Date]', '[Recipient name and organisation]', '[Address or contact]']),
        ('Message', ['Dear [recipient],', '[Introduce the purpose of your letter.]',
                     '[Add the key details and the action you are requesting.]',
                     'Thank you for your time.']),
        ('Closing', ['Kind regards,', '[Your name]  |  [Your role]']),
    ])
    word_template(OUT / 'project-brief.docx', 'Project brief',
                  'A structured fictional brief for planning a project', [
        ('Overview', ['[Summarise the project and its purpose in two sentences.]']),
        ('Objectives', ['[Objective 1]', '[Objective 2]', '[Objective 3]']),
        ('Scope and audience', ['[Describe what is included, who it serves and what is out of scope.]']),
        ('Milestones', ['[Milestone] — [owner] — [target date]', '[Milestone] — [owner] — [target date]']),
        ('Notes', ['[Add assumptions, risks and next actions.]']),
    ])
    spreadsheet(OUT / 'monthly-budget.xlsx', title='Monthly budget',
                subtitle='Fictional amounts; replace values and confirm calculations.',
                columns=[('Category', 24), ('Planned', 17), ('Actual', 17), ('Variance', 17), ('Notes', 35)],
                rows=[('Materials', 25000, 23000, '=C5-B5', 'Example only'),
                      ('Transport', 15000, 16000, '=C6-B6', 'Example only'),
                      ('Other', 10000, 8000, '=C7-B7', 'Example only')],
                formula_column={2, 3, 4}, notes=[
                    'MONTHLY BUDGET — fictional starter',
                    'Replace the sample categories and amounts. Check all formulas in your spreadsheet app.',
                    'DeskSuite content editing/export does not preserve formulas, styles or exact layout.',
                ])
    spreadsheet(OUT / 'stock-tracker.xlsx', title='Stock tracker',
                subtitle='Fictional rows and values; review before using for real records.',
                columns=[('Item', 28), ('Quantity', 15), ('Unit cost', 17), ('Value', 17), ('Note', 35)],
                rows=[('Paper packs', 12, 4500, '=B5*C5', 'Example only'),
                      ('Folders', 20, 900, '=B6*C6', 'Example only'),
                      ('Pens', 15, 600, '=B7*C7', 'Example only')],
                formula_column={3, 4}, notes=[
                    'STOCK TRACKER — fictional starter',
                    'Replace items and amounts. Check all formulas in your spreadsheet app.',
                    'DeskSuite content editing/export does not preserve formulas, styles or exact layout.',
                ])
    make_presentation(OUT / 'project-pitch.pptx', 'Project pitch', [
        cover('Your project\npitch', 'A clear introduction to your idea', 'Project pitch starter'),
        section_slide('The opportunity', '01  /  Context', [
            ('The need', 'Describe the problem\nand who experiences it.'),
            ('The idea', 'Summarise your proposed\nsolution in plain language.'),
            ('The impact', 'Describe the change\nyou want to make.'),
        ]),
        section_slide('The approach', '02  /  Plan', [
            ('Discover', 'Listen, research and\nvalidate the need.'),
            ('Build', 'Describe the first\npractical steps.'),
            ('Measure', 'Define useful outcomes\nand how to review them.'),
        ]),
        section_slide('Next steps', '03  /  Action', [
            ('Priority', 'Choose one action\nto start with.'),
            ('People', 'Name the roles and\nresources needed.'),
            ('Timeline', 'Add an achievable\nreview date.'),
        ]),
    ])
    make_presentation(OUT / 'workshop-lesson.pptx', 'Workshop lesson', [
        cover('Your workshop\nsession', 'An editable learning deck for a fictional session', 'Workshop starter'),
        section_slide('What we will learn', '01  /  Objectives', [
            ('Understand', 'Introduce the topic\nand why it matters.'),
            ('Practise', 'Describe a hands-on\nactivity for learners.'),
            ('Review', 'Explain how learners\ncan check their work.'),
        ]),
        section_slide('Try it together', '02  /  Activity', [
            ('Prepare', 'List the materials\nand starting point.'),
            ('Create', 'Write simple steps\nfor the exercise.'),
            ('Reflect', 'Add questions for\nreview and discussion.'),
        ]),
    ])
    for path in sorted(OUT.iterdir()):
        if path.is_file():
            print(f'{path.relative_to(ROOT)}: {path.stat().st_size:,} bytes')


if __name__ == '__main__':
    main()
