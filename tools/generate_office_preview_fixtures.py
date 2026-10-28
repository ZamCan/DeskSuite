#!/usr/bin/env python3
"""Create fictional Office files for local renderer tests (not user data).

Requires python-docx, openpyxl, python-pptx and xlwt, installed by the developer
in their existing Python environment. These are test-generation dependencies,
not runtime app dependencies. No legacy DOC/PPT writer is used here.
"""
from datetime import datetime
from pathlib import Path

from docx import Document
from openpyxl import Workbook as ExcelWorkbook
from pptx import Presentation
import xlwt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'test' / 'fixtures' / 'office_preview'
STAMP = datetime(2026, 1, 1, 0, 0, 0)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.core_properties.title = 'DeskSuite preview fixture'
    doc.core_properties.author = 'DeskSuite test'
    doc.core_properties.created = STAMP
    doc.core_properties.modified = STAMP
    doc.add_heading('DeskSuite Preview', level=1)
    doc.add_paragraph('Offline document rendering check.')
    table = doc.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [('Item', 'Quantity'), ('Paper', '12')]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    doc.save(OUT / 'preview.docx')

    workbook = ExcelWorkbook()
    workbook.properties.creator = 'DeskSuite test'
    workbook.properties.created = STAMP
    workbook.properties.modified = STAMP
    sheet = workbook.active
    sheet.title = 'Inventory'
    for row in [('DeskSuite Preview', 'Quantity'), ('Paper', 12)]:
        sheet.append(row)
    other = workbook.create_sheet('Sales')
    other.append(('Local sample', 1))
    workbook.save(OUT / 'preview.xlsx')

    presentation = Presentation()
    presentation.core_properties.author = 'DeskSuite test'
    presentation.core_properties.title = 'DeskSuite preview fixture'
    presentation.core_properties.created = STAMP
    presentation.core_properties.modified = STAMP
    first = presentation.slides.add_slide(presentation.slide_layouts[1])
    first.shapes.title.text = 'DeskSuite Preview'
    first.placeholders[1].text = 'Offline rendering check.'
    second = presentation.slides.add_slide(presentation.slide_layouts[1])
    second.shapes.title.text = 'Second slide'
    second.placeholders[1].text = 'A fictional sample.'
    presentation.save(OUT / 'preview.pptx')

    legacy = xlwt.Workbook()
    inventory = legacy.add_sheet('Inventory')
    for y, row in enumerate([('DeskSuite Preview', 'Quantity'), ('Paper', 12)]):
        for x, value in enumerate(row):
            inventory.write(y, x, value)
    sales = legacy.add_sheet('Sales')
    sales.write(0, 0, 'Local sample')
    legacy.save(str(OUT / 'preview.xls'))
    print(f'Created four fictional Office fixtures in {OUT.relative_to(ROOT)}')


if __name__ == '__main__':
    main()
